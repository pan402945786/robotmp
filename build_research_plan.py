from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(r'D:\ww2\robotmp')
OUT = ROOT / 'research_plan_predictive_robotics_skills.docx'
MD = ROOT / 'research_plan_predictive_robotics_skills.md'

# ---------- Content ----------
content = {
    'title': '基于预测误差与失败经验 Skill 的机器人自主学习研究计划',
    'subtitle': '面向机器人操作任务的预测 观测 诊断 验证与经验复用闭环',
    'date': '2026 年 10 月 17 日',
    'project': 'D:\\ww2\\robotmp MVP',
}

sections = [
('一 研究定位与摘要', [
('本研究拟解决的问题', '本研究关注机器人在执行抓取、搬运和放置任务时，如何把一次失败转化为下一次可复用的经验。研究对象不是单纯的故障报警，也不是让大模型直接控制关节，而是建立一条可验证的闭环：机器人先对动作结果作出结构化预测，执行动作后采集实际状态，比较预测与现实的偏差，诊断偏差原因，生成候选失败经验 Skill，在仿真或安全环境中验证 Skill，最后把已验证的 Skill 用于后续任务的预测和恢复。'),
('当前判断', '现有机器人系统已经具备感知、规划、模仿学习、局部恢复和技能复用等能力，但公开可用的工业系统通常还没有把“失败诊断、经验验证、版本管理和跨任务复用”整合成一个长期运行的自主学习闭环。本项目的研究价值在于把失败经验从自然语言描述提升为带触发条件、证据链、适用范围、验证结果和回滚机制的结构化对象。'),
('本计划的用途', '本文档用于指导后续的系统实现、实验设计、数据记录、论文撰写和阶段评审。它明确区分当前 MVP 已经验证的工程事实、尚待验证的科学假设，以及每个阶段必须交付的证据。'),
]),
('二 科学问题与研究假设', [
('核心科学问题', '在存在重复故障、传感器不确定性和环境变化的机器人操作任务中，经过证据验证的失败经验 Skill，是否能够降低后续任务的错误成功判断，提高失败恢复率，并迁移到未见过的相似任务？'),
('研究假设 H1', '与固定计划系统相比，加入预测与实际状态比较的反馈系统能够显著降低 false-success rate，并提高任务成功率和失败恢复率。'),
('研究假设 H2', '与只进行局部反馈恢复但不保存经验的系统相比，加入经过验证的失败经验 Skill 后，系统在重复故障和相似新场景中的恢复率更高、重复失败率更低。'),
('研究假设 H3', '将事实、原因假设、预测修正和恢复动作分开记录，并要求 Skill 经过独立验证，可以在不增加不安全动作的前提下提高跨物体、跨姿态和跨环境的迁移效果。'),
('可证伪条件', '如果在严格的独立测试集上，Skill 系统相对于反馈基线没有稳定提升，或者提升只出现在训练过的故障条件中，或者安全违规率显著上升，则上述假设不成立或需要缩小适用范围。'),
]),
('三 研究目标与预期贡献', [
('目标一 建立可审计的闭环原型', '把当前 robotmp MVP 扩展为可记录 episode、预测、观测、偏差、诊断、恢复和 Skill 版本的研究平台。'),
('目标二 建立失败经验表示', '定义统一的 Skill 数据结构，明确触发条件、事实、待验证假设、预测修正、恢复动作、适用范围、验证指标、版本和回滚状态。'),
('目标三 验证 Skill 的实际作用', '通过 A、B、C 三组对照实验，区分反馈恢复本身的作用和经验复用带来的额外收益。'),
('目标四 研究迁移与安全边界', '测试 Skill 在新物体、新抓取姿态、新噪声水平和新故障组合下的表现，测量是否引入新的不安全动作。'),
('预期贡献', '形成一套适用于机器人操作任务的预测误差与失败经验闭环架构；提出可验证、可回滚的 Skill 生命周期机制；建立覆盖任务成功、错误成功判断、恢复能力、迁移能力和安全性的评价协议；完成一个可复现实验基线和研究数据集。'),
]),
('四 研究范围与边界', [
('首个研究任务', '以桌面抓取立方体或杯子并放置到托盘为主要任务，动作集合包括观察、接近、闭合夹爪、抬升、搬运、释放、重新观察、刷新位姿、重试抓取和安全停止。'),
('首批故障类型', '抓取点偏移、夹持力不足、物体位姿过期、物体滑落、视觉短时遮挡、意外接触和碰撞风险。第一阶段只要求稳定覆盖抓取失败未抬升这一主故障，后续逐步加入其他故障。'),
('明确边界', '大模型只负责任务级规划、异常解释和候选 Skill 生成；高频控制、碰撞检测、力矩限制、急停和动作许可由本地控制器与安全层负责。一次失败只能产生 candidate Skill，不能直接修改生产策略。'),
]),
('五 系统总体架构', [
('数据与状态层', '保留图像、深度、关节、末端位姿、力和夹爪信号等原始数据，同时构建结构化 SceneState，记录对象位置、可见性、接触、抓取、抬升、置信度、阶段和事件。自然语言只作为解释接口，不作为唯一事实载体。'),
('预测与监控层', '每个关键动作执行前生成 Prediction，包含预期阶段、预期事件、连续状态变化、置信度、容差和适用 Skill。动作完成后由比较器输出 match、uncertain 或 deviation。证据不足时先重新观察，不直接生成失败经验。'),
('诊断与恢复层', '诊断器把偏差拆成可观察事实、可能原因和待验证假设，再从经过许可的恢复动作库中选择方案。恢复动作必须满足前置条件、参数范围和安全约束。'),
('Skill 生命周期层', '候选 Skill 进入 candidate 状态，经过仿真和安全实验后进入 testing，再根据预设门槛升级为 verified。表现下降、适用范围超出或与其他规则冲突时进入 disabled，并支持回滚。'),
('闭环流程', '感知物理世界 → 构建结构化状态 → 检索相关 Skill → 预测未来状态 → 执行动作 → 采集实际观测 → 比较预测与现实 → 继续、重观察或诊断恢复 → 生成候选 Skill → 验证并更新经验库。'),
]),
('六 关键数据结构与概念定义', [
('Episode', '一次完整任务的可审计记录，至少包含初始状态、动作序列、预测、观测、比较结果、诊断、恢复过程、最终结果和时间戳。'),
('Prediction', '对未来状态或事件的结构化预测，不仅包含位置和阶段，还要包含预期事件、置信度、容差和所应用的 Skill。'),
('Comparison', '预测与实际观测之间的比较结果。match 表示证据支持继续执行；uncertain 表示证据不足；deviation 表示有足够证据表明预测偏离现实。'),
('Skill', '带有触发条件、观察证据、事实、原因假设、预测修正、恢复动作、适用范围、验证指标和版本信息的经验对象。'),
('事实与假设的区分', '例如“夹爪已闭合”和“物体没有抬升”是观测或状态判断；“抓取点偏移”和“夹持力不足”是待验证原因，不能在第一次失败后直接当作事实。'),
]),
('七 实验设计', [
('对照组定义', 'A 组为固定计划系统：执行预先定义的动作序列，不进行预测与现实比较。B 组为反馈恢复系统：比较预测与观测，发现偏差后执行局部恢复，但不保存和检索跨任务 Skill。C 组为反馈恢复加 Skill 系统：在 B 组基础上生成候选 Skill，经过独立验证后在后续任务开始前检索和应用。'),
('训练与测试划分', 'Skill 只能从训练场景产生。测试场景必须使用未见过的物体实例、抓取姿态、噪声参数或故障组合。禁止把测试结果用于即时升级 Skill，避免数据泄漏。'),
('场景变量', '物体类型和尺寸、初始位姿、抓取点偏移、摩擦力、夹持力、视觉遮挡比例、传感器噪声、观测延迟、动作执行误差和故障发生阶段。'),
('首轮实验矩阵', '第一轮固定任务为抓取并放置，故障包含正常、抓取点偏移、夹持力不足、位姿过期和视觉短暂遮挡。每种故障至少设置多个随机种子和多个物体姿态；训练和测试使用不同随机种子。'),
('重复次数建议', '探索性实验每个条件至少 30 次；正式对照实验每个条件建议 100 次以上。所有 trial 保存随机种子、完整轨迹和系统版本。若资源有限，应优先保证独立测试集和跨场景测试，而不是只增加同一确定性场景的重复次数。'),
]),
('八 评价指标与统计方法', [
('主要指标', '任务成功率、false-success rate、失败恢复率、重复同类失败频率、跨场景迁移成功率和不安全动作次数。'),
('效率指标', '平均完成时间、恢复耗时、重新观察次数、动作执行次数、模型调用次数和推理延迟。'),
('质量指标', '失败诊断与真实注入原因的一致性、Skill 触发精确率、Skill 误用率、候选 Skill 验证通过率、Skill 版本回滚次数和证据完整率。'),
('安全指标', '碰撞次数、超力矩次数、越界动作次数、急停次数、人工接管次数和恢复过程中对环境造成的额外损害。安全违规应作为硬约束，而不是与成功率简单加权平均。'),
('统计分析', '报告每个条件的均值、标准差或置信区间，并使用按场景或随机种子分层的统计检验。除了总体成功率，还要报告不同故障类型和未见场景上的分组结果。若样本量不足，应明确报告不确定性，而不是只给单一百分比。'),
]),
('九 Skill 验证与安全策略', [
('升级门槛', '候选 Skill 至少经过预设数量的独立试验；恢复成功率达到目标阈值；false-success rate 不高于阈值；安全违规次数为零；适用范围没有超出验证范围；与已部署 Skill 不发生未处理的冲突。'),
('建议初始策略', '最少 20 次验证用于 MVP，正式研究建议跨多个场景累计 100 次以上；恢复成功率目标不低于 0.70；false-success rate 目标不高于 0.10；最大安全违规次数为 0。具体阈值应在试验前锁定。'),
('灰度发布', 'verified Skill 先进入 shadow 模式，只记录如果应用它会采取什么策略，不直接影响动作；连续回放和安全试验通过后进入小比例部署，再根据在线结果决定是否扩大范围。'),
('回滚机制', '每个 Skill 保留版本、来源 episode、验证数据、部署范围和停用原因。新版本表现下降或触发冲突时，恢复到上一个 verified 版本。'),
]),
('十 实施路线图', [
('阶段一 研究基础整理 1 至 2 周', '锁定问题定义、假设、指标、故障分类、数据字典和实验协议；为当前 robotmp 建立版本标签和可重复运行入口。'),
('阶段二 可审计日志 2 至 4 周', '新增 EpisodeStore、事件日志、预测记录、观测记录、比较结果和实验元数据；将 JSON 记录逐步迁移到 SQLite。'),
('阶段三 多故障仿真 4 至 8 周', '扩展故障注入、随机化环境和独立测试集；让恢复动作真实改变环境状态，避免通过简单地清除故障来制造恢复成功。'),
('阶段四 Skill 生命周期 8 至 12 周', '实现 candidate、testing、verified、shadow、deployed、disabled 状态，增加验证门槛、冲突检查、版本控制和回滚。'),
('阶段五 对照实验 12 至 16 周', '完成 A/B/C 实验、消融实验和迁移实验，生成图表、统计结果和失败案例分析。'),
('阶段六 真实机器人小规模验证 16 周以后', '先在低风险、低速度、带急停和人工监督的任务上验证状态估计、失败检测和恢复动作，再逐步扩大任务范围。'),
]),
('十一 当前 MVP 到研究平台的改造清单', [
('已有基础', '当前 robotmp 已实现轻量环境、SceneState、Prediction、预测比较、规则诊断、JSON SkillStore、Skill 验证和 baseline、feedback、skills 三种运行模式。'),
('优先补齐', '完整 Episode 记录、更多故障模式、随机化、独立训练测试划分、真实恢复效果、可回滚 Skill 版本、统计分析脚本和安全指标。'),
('需要修正的实验假设', '当前模拟环境中的 retry_grasp 会直接把故障模式清除，因此只能验证控制流，不能证明系统学会了真实恢复。下一版应让抓取点、夹持力或位姿刷新真正改变成功概率。'),
('暂不优先', '不急于训练端到端大模型，不让大模型直接输出关节指令，也不在缺乏证据时把一次失败原因写成确定事实。先保证数据、基线和实验可重复。'),
]),
('十二 风险与应对', [
('学习到错误经验', '要求事实和假设分离；候选 Skill 必须经过独立验证；记录来源证据；支持禁用和回滚。'),
('训练测试泄漏', '按场景、随机种子和故障组合划分数据；测试期间冻结 Skill 库；实验结束后再统一分析。'),
('只改善模拟器', '采用域随机化和未见场景测试；最终在真实设备或更高保真仿真中验证关键结论。'),
('大模型产生不可执行建议', '使用 JSON Schema、动作白名单、参数范围、行为树和本地安全层校验；拒绝无法验证或越界的动作。'),
('指标被单一成功率掩盖', '同时报告错误成功、恢复、安全、时间、成本和迁移指标；把安全违规作为硬约束。'),
]),
('十三 阶段性交付物', [
('代码交付', '可重复运行的环境、状态估计器、预测器、比较器、诊断器、恢复策略库、SkillStore、验证器和实验脚本。'),
('数据交付', '完整 episode 日志、故障注入配置、随机种子、训练测试清单、Skill 版本记录和失败案例集。'),
('研究交付', '实验协议、A/B/C 结果、消融实验、迁移实验、统计分析、系统限制和下一阶段问题清单。'),
('论文材料', '问题定义、方法图、数据结构、实验表格、失败案例可视化、结果讨论和可复现说明。'),
]),
('十四 首个 30 天执行计划', [
('第 1 周', '确定研究问题、H1 至 H3、指标定义和首轮故障矩阵；冻结当前 MVP 版本作为 baseline。'),
('第 2 周', '实现 episode 级日志和 SQLite 存储；为每次预测、观测、比较和恢复写入时间戳与版本号。'),
('第 3 周', '加入抓取点偏移、夹持力变化、位姿过期和视觉遮挡四类故障；实现随机种子和独立测试场景。'),
('第 4 周', '完成第一轮 A/B/C 实验和可视化；检查 Skill 是否真正改变后续预测或恢复策略，形成阶段报告。'),
]),
('十五 研究记录模板', [
('每次实验应记录', '实验日期、代码版本、环境版本、机器人或仿真器配置、任务类型、物体类型、故障类型、随机种子、预测、观测、偏差、诊断、恢复动作、最终结果、安全事件和备注。'),
('每个 Skill 应记录', 'Skill ID、版本、来源 episode、触发条件、证据、事实、原因假设、预测修正、恢复动作、适用范围、验证试验、统计结果、部署状态、冲突记录和回滚原因。'),
('阶段评审问题', '是否有独立测试集？是否能复现？Skill 是否改善了后续任务？改善是否来自 Skill 而非恢复代码？是否发生安全退化？结论适用到什么范围？'),
]),
('十六 预期结论形式', [
('如果结果支持假设', '可以表述为：在规定的仿真任务、故障类型和安全约束下，经过验证的失败经验 Skill 相比固定计划和只反馈恢复基线，降低了错误成功判断并提高了后续恢复或迁移表现。'),
('如果结果不支持假设', '应明确指出 Skill 没有带来可重复收益，或收益依赖特定故障、特定物体或特定仿真器，并据此缩小研究主张或改进状态表示、验证机制和恢复策略。'),
('研究边界声明', '本研究不以证明机器人具备一般智能为目标，而是验证一个可审计、可验证、可回滚的失败经验闭环是否能改善特定机器人操作任务。'),
]),
]

# ---------- Markdown ----------
md_lines = [f"# {content['title']}", '', content['subtitle'], '', f"编制日期：{content['date']}", f"项目基础：`{content['project']}`", '']
for h, items in sections:
    md_lines.append(f"## {h}")
    md_lines.append('')
    for sh, body in items:
        md_lines.append(f"### {sh}")
        md_lines.append('')
        md_lines.append(body)
        md_lines.append('')
MD.write_text('\n'.join(md_lines), encoding='utf-8')

# ---------- DOCX helpers ----------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_borders(table, color='D9D9D9', size='6'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def set_run_font(run, name='Aptos', size=10.5, bold=False, color='000000'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style='Body Text')
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p

# ---------- DOCX build ----------
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
styles['Normal'].font.size = Pt(10.5)

for style_name, size, color in [('Title', 22, '000000'), ('Heading 1', 15, '000000'), ('Heading 2', 12, '000000'), ('Heading 3', 10.5, '000000')]:
    st = styles[style_name]
    st.font.name = 'Aptos Display' if style_name in ('Title','Heading 1') else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(10 if style_name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)

# custom body style
if 'Body Text' not in styles:
    body_style = styles.add_style('Body Text', WD_STYLE_TYPE.PARAGRAPH)
else:
    body_style = styles['Body Text']
body_style.font.name = 'Aptos'
body_style._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
body_style.font.size = Pt(10.5)
body_style.paragraph_format.first_line_indent = Inches(0.25)
body_style.paragraph_format.line_spacing = 1.18

# Header/footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rh = header.add_run('robotmp 研究计划')
set_run_font(rh, size=8.5, color='666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = footer.add_run('页码 ')
set_run_font(rf, size=8.5, color='666666')
add_page_number(footer)

# Title page
title_p = doc.add_paragraph(style='Title')
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(45)
title_p.paragraph_format.space_after = Pt(16)
r = title_p.add_run(content['title'])
set_run_font(r, name='Aptos Display', size=22, bold=True)
sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(28)
r = sub_p.add_run(content['subtitle'])
set_run_font(r, size=12, color='404040')

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
meta.columns[0].width = Inches(1.35)
meta.columns[1].width = Inches(4.7)
set_table_borders(meta)
meta_rows = [('文档用途', '用于指导后续系统实现、实验设计、数据记录和论文准备'), ('项目基础', r'D:\ww2\robotmp MVP'), ('编制日期', content['date'])]
for i, (a,b) in enumerate(meta_rows):
    for j, txt in enumerate((a,b)):
        cell = meta.cell(i,j)
        cell.width = meta.columns[j].width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if j == 0:
            set_cell_shading(cell, 'EAF2F8')
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rr = p.add_run(txt)
        set_run_font(rr, size=9.5, bold=(j==0))

doc.add_paragraph('')
intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(30)
intro.paragraph_format.space_after = Pt(8)
r = intro.add_run('研究计划摘要')
set_run_font(r, size=12, bold=True)
add_body(doc, '本计划围绕机器人在操作任务中“预测未来状态、执行动作、比较偏差、诊断失败、验证经验并复用 Skill”的闭环展开。当前项目已经具备可运行 MVP，下一阶段的重点是把演示型程序推进为可审计、可重复、可统计分析的研究平台，并用 A、B、C 三组对照实验验证失败经验 Skill 是否确实改善后续任务表现。')
add_body(doc, '研究主张将被限定在明确的任务、故障类型和安全边界内。一次失败只产生候选 Skill，只有经过独立验证且没有增加安全风险的经验才允许进入 verified 状态。')

doc.add_page_break()

# Content sections
for h, items in sections:
    h1 = doc.add_paragraph(style='Heading 1')
    rr = h1.add_run(h)
    set_run_font(rr, size=15, bold=True)
    for sh, body in items:
        h2 = doc.add_paragraph(style='Heading 2')
        rr = h2.add_run(sh)
        set_run_font(rr, size=12, bold=True)
        add_body(doc, body)
    # Add compact separator via spacing only

# Final review checklist table
doc.add_page_break()
h = doc.add_paragraph(style='Heading 1')
r = h.add_run('附录 研究启动检查表')
set_run_font(r, size=15, bold=True)
check_rows = [
    ('问题', '是否已经把科学问题写成可证伪的形式？', '□'),
    ('基线', 'A、B、C 三组系统的行为是否已经固定并可重复运行？', '□'),
    ('数据', '每次 trial 是否保存随机种子、完整轨迹和版本信息？', '□'),
    ('独立性', '测试场景是否与 Skill 产生场景分离？', '□'),
    ('安全', '是否定义碰撞、越界、超力矩和人工接管等硬约束？', '□'),
    ('验证', 'candidate Skill 是否经过独立验证后才可升级？', '□'),
    ('解释', '事实、假设、恢复动作和最终结论是否分开记录？', '□'),
    ('复现', '其他人是否可以仅凭代码、配置和数据重现实验？', '□'),
]
t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [Inches(0.8), Inches(5.5), Inches(0.45)]
for i,w in enumerate(widths): t.columns[i].width = w
set_table_borders(t)
headers = ['类别', '检查项', '完成']
for j,txt in enumerate(headers):
    cell=t.cell(0,j); cell.width=widths[j]; set_cell_shading(cell,'1F4E79'); set_cell_margins(cell)
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(txt); set_run_font(r,size=9.5,bold=True,color='FFFFFF')
set_repeat_table_header(t.rows[0])
for i,row in enumerate(check_rows, start=1):
    cells=t.add_row().cells
    for j,txt in enumerate(row):
        cells[j].width=widths[j]; set_cell_margins(cells[j]); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i%2==0: set_cell_shading(cells[j],'F5F8FB')
        p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=1 else WD_ALIGN_PARAGRAPH.LEFT
        r=p.add_run(txt); set_run_font(r,size=9.5)

# Set all table alignment and add keep lines
for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.08

# core properties
props = doc.core_properties
props.title = content['title']
props.subject = '机器人自主学习与失败经验 Skill 研究计划'
props.author = 'robotmp research'
props.keywords = 'robotics, failure recovery, skill learning, predictive monitoring'

doc.save(OUT)
print(OUT)
print(MD)

